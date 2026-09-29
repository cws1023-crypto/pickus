# -*- coding: utf-8 -*-
"""
강남역 반경 N km 이내 병·의원 수집기

데이터 출처
  1) 건강보험심사평가원 병원정보서비스 (공공데이터포털, 무료 인증키 필요)
     - 병원명, 대표전화, 주소, 홈페이지, 종별(업종), 거리
  2) 병원 홈페이지 공개 이메일 (선택)
     - 홈페이지에 '이메일 무단수집 거부'가 게시된 곳은 수집하지 않음 (정보통신망법 제50조의2)

사용법
  python collect_hospitals.py --key 공공데이터포털_인증키
  python collect_hospitals.py --key 인증키 --radius 3000 --no-email
"""
import argparse
import datetime as dt
import os
import re
import sys
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from urllib.parse import urljoin, urlparse

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# 강남역(2호선) 좌표 (WGS84)
GANGNAM_X = 127.027621
GANGNAM_Y = 37.497942

HIRA_URL = "https://apis.data.go.kr/B551182/hospInfoServicev2/getHospBasisList"
DEFAULT_OUT_DIR = r"C:\Users\user\OneDrive\Desktop\오파독 작업\병원자동수집"

HEADERS = [
    "수집일자", "수집경로", "수집경로세부", "이름", "회사명", "직책", "휴대폰",
    "회사전화", "팩스", "이메일", "주소", "홈페이지", "업종",
    "강남역거리(m)", "의사수", "개설일자", "이메일수집메모",
]

UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                    "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"}

EMAIL_RE = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")
REFUSAL_RE = re.compile(r"(이메일|전자우편|e-?mail)[^<]{0,30}무단\s*수집|무단\s*수집\s*거부", re.I)
BAD_EMAIL_SUFFIX = (".png", ".jpg", ".jpeg", ".gif", ".webp", ".svg", ".css", ".js")
BAD_EMAIL_DOMAINS = ("example.com", "sentry.io", "wixpress.com", "domain.com", "email.com")
CONTACT_HINTS = ("contact", "about", "intro", "info", "map", "location", "privacy",
                 "오시는", "문의", "소개", "개인정보", "찾아오")


# ---------------------------------------------------------------- 심평원 API
def fetch_hira(key, radius, rows=1000):
    items, page = [], 1
    while True:
        params = {
            "serviceKey": key, "pageNo": page, "numOfRows": rows, "_type": "json",
            "xPos": GANGNAM_X, "yPos": GANGNAM_Y, "radius": radius,
        }
        for attempt in range(4):
            try:
                r = requests.get(HIRA_URL, params=params, timeout=30)
                r.raise_for_status()
                data = r.json()
                break
            except ValueError:
                sys.exit(f"[오류] API 응답이 JSON이 아닙니다. 인증키/활용신청 상태를 확인하세요.\n{r.text[:300]}")
            except requests.RequestException as e:
                if attempt == 3:
                    raise
                print(f"  재시도({attempt + 1}) - {e}")
                time.sleep(2 ** attempt)

        body = data.get("response", {}).get("body", {})
        header = data.get("response", {}).get("header", {})
        if header.get("resultCode") not in ("00", None):
            sys.exit(f"[오류] API: {header}")
        got = (body.get("items") or {}).get("item") or []
        if isinstance(got, dict):
            got = [got]
        items.extend(got)
        total = int(body.get("totalCount", 0))
        print(f"  {len(items)}/{total} 건 수신")
        if len(items) >= total or not got:
            return items
        page += 1


# ---------------------------------------------------------------- 이메일 수집
def normalize_url(u):
    u = (u or "").strip()
    if not u or u in ("-", "http://", "https://"):
        return ""
    if not u.startswith(("http://", "https://")):
        u = "http://" + u
    return u


def get_html(url):
    r = requests.get(url, headers=UA, timeout=10, allow_redirects=True)
    r.encoding = r.apparent_encoding or r.encoding
    return r.text, r.url


def extract_emails(html):
    found = set()
    for m in EMAIL_RE.findall(html):
        e = m.strip(".").lower()
        if e.endswith(BAD_EMAIL_SUFFIX) or e.split("@")[1] in BAD_EMAIL_DOMAINS:
            continue
        found.add(e)
    return found


def crawl_email(homepage):
    """(이메일들, 메모) 반환"""
    url = normalize_url(homepage)
    if not url:
        return "", "홈페이지 없음"
    try:
        html, final = get_html(url)
    except requests.RequestException:
        return "", "홈페이지 접속불가"

    pages = [html]
    host = urlparse(final).netloc
    links = re.findall(r'href=["\']([^"\'#]+)["\']', html, re.I)
    sub = []
    for href in links:
        full = urljoin(final, href)
        if urlparse(full).netloc == host and any(h in full.lower() for h in CONTACT_HINTS):
            if full not in sub:
                sub.append(full)
    for s in sub[:3]:
        try:
            pages.append(get_html(s)[0])
        except requests.RequestException:
            pass

    if any(REFUSAL_RE.search(p) for p in pages):
        return "", "무단수집거부 게시 - 미수집"
    emails = set()
    for p in pages:
        emails |= extract_emails(p)
    if not emails:
        return "", "공개 이메일 없음"
    return ", ".join(sorted(emails)), "홈페이지 공개 이메일"


# ---------------------------------------------------------------- 엑셀 저장
def save_xlsx(rows, path, radius_km):
    wb = Workbook()
    ws = wb.active
    ws.title = f"강남역 {radius_km:g}km 병의원"
    ws.append(HEADERS)
    for r in rows:
        ws.append([r.get(h, "") for h in HEADERS])

    head_fill = PatternFill("solid", fgColor="D9EAD3")
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
    widths = [11, 16, 16, 8, 30, 8, 14, 14, 14, 30, 50, 30, 12, 12, 8, 11, 22]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def main():
    ap = argparse.ArgumentParser(description="강남역 주변 병의원 수집")
    ap.add_argument("--key", default=os.environ.get("DATA_GO_KR_KEY"),
                    help="공공데이터포털 인증키(Decoding 키 권장)")
    ap.add_argument("--radius", type=int, default=3000, help="반경(m), 기본 3000")
    ap.add_argument("--out", default=DEFAULT_OUT_DIR, help="저장 폴더")
    ap.add_argument("--no-email", action="store_true", help="홈페이지 이메일 수집 생략")
    ap.add_argument("--workers", type=int, default=16)
    a = ap.parse_args()
    if not a.key:
        sys.exit("인증키가 필요합니다: --key 인증키  (README.md 참고)")

    print(f"[1/3] 심평원 병원정보 조회 (강남역 반경 {a.radius}m)")
    items = fetch_hira(a.key, a.radius)

    today = dt.date.today().strftime("%Y-%m-%d")
    rows = []
    for it in items:
        rows.append({
            "수집일자": today,
            "수집경로": "심평원 병원정보",
            "수집경로세부": f"강남역 반경{a.radius / 1000:g}km",
            "회사명": it.get("yadmNm", ""),
            "회사전화": it.get("telno", ""),
            "주소": it.get("addr", ""),
            "홈페이지": normalize_url(it.get("hospUrl", "")),
            "업종": it.get("clCdNm", ""),
            "강남역거리(m)": round(float(it.get("distance") or 0)),
            "의사수": it.get("drTotCnt", ""),
            "개설일자": str(it.get("estbDd", "")),
        })
    rows.sort(key=lambda r: r["강남역거리(m)"])

    if a.no_email:
        print("[2/3] 이메일 수집 생략")
    else:
        targets = [r for r in rows if r["홈페이지"]]
        print(f"[2/3] 홈페이지 이메일 수집 ({len(targets)}곳)")
        with ThreadPoolExecutor(a.workers) as ex:
            futs = {ex.submit(crawl_email, r["홈페이지"]): r for r in targets}
            for n, f in enumerate(as_completed(futs), 1):
                r = futs[f]
                try:
                    r["이메일"], r["이메일수집메모"] = f.result()
                except Exception as e:  # 개별 사이트 오류는 건너뜀
                    r["이메일수집메모"] = f"오류: {type(e).__name__}"
                if n % 50 == 0:
                    print(f"  {n}/{len(targets)}")
        for r in rows:
            r.setdefault("이메일수집메모", "홈페이지 없음")

    os.makedirs(a.out, exist_ok=True)
    fname = f"강남역_{a.radius / 1000:g}km_병의원_{dt.date.today():%Y%m%d}.xlsx"
    path = os.path.join(a.out, fname)
    save_xlsx(rows, path, a.radius / 1000)
    n_mail = sum(1 for r in rows if r.get("이메일"))
    print(f"[3/3] 완료: {len(rows)}곳 (이메일 {n_mail}곳)\n  -> {path}")


if __name__ == "__main__":
    main()
